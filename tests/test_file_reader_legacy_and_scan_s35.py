"""S35 席位锁：旧格式（.xls/.ppt/.doc）与扫描件 PDF 的「读不了要诚实说」两条腿。

钉三件事（写法对齐 ``test_file_reader_parse_safety`` / ``test_pdf_scan_honesty``）：

① **旧格式必进诚实态、不得冒充「内容为空」**——真 OLE2 件必须交回
   ``parser_unavailable``（V2.1 风险 8 那三条锁的语义原样保留），且
   ``file_read_failure_note`` 要说「读不出来」并带上**本机现算**的缺件代号；
   后缀写旧格式、内容实为 OOXML 的件按魔数改道（能真读就真读，读不出也不许
   再谎称它是旧 OLE2）。
② **新接入腿在缺能力机器上必降级不崩**——把 ``shutil.which`` /
   ``importlib.util.find_spec`` / ``winreg`` 三样探测全打成「什么都没有」，
   读取链仍照常返回诚实降级、不抛异常、不静默。
③ **文本白名单与 120000 截断零回归**——本席只加识别，没动那两把尺
   （白名单里绝不允许出现 Office/PDF 后缀，否则文本腿会把容器当纯文本吐出去）。

样本全在 ``tmp_path``（``--basetemp`` 指仓库外），源码树不留件。
"""

from __future__ import annotations

import importlib.util
import shutil
from pathlib import Path
from typing import Any

import pytest

from plugins.bot_unified_runtime.domains.files.sources import file_reader

_OLE2 = bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 600


def _poison_probes(monkeypatch: pytest.MonkeyPatch) -> None:
    """把三样本机探测全打成「缺件」（等价于换一台什么都没装的机器）。"""

    def _no_tool(_name: str) -> None:
        return None

    def _no_spec(_name: str) -> None:
        return None

    def _no_key(*_a: Any, **_k: Any) -> None:
        raise OSError("no registry here")

    monkeypatch.setattr(shutil, "which", _no_tool)
    monkeypatch.setattr(importlib.util, "find_spec", _no_spec)
    try:
        import winreg

        monkeypatch.setattr(winreg, "OpenKey", _no_key)
    except ImportError:  # pragma: no cover - 非 Windows 上本就没有这一格
        pass


# --------------------------------------------------------------------------- ①


def test_real_ole2_legacy_files_stay_honest(tmp_path: Path) -> None:
    """.xls/.ppt/.doc 三形确为 OLE2 ⇒ parser_unavailable + 正文空（绝不冒充「读到了空文档」）。"""
    for name, kind, phrase in (
        ("old.xls", "spreadsheet", "旧版 Excel 二进制"),
        ("old.ppt", "presentation", "旧版 PowerPoint 二进制"),
        ("old.doc", "document", "旧版 Word 二进制"),
    ):
        target = tmp_path / name
        target.write_bytes(_OLE2)
        result = file_reader.read_supported_file(target)
        assert result.kind == kind
        assert result.metadata
        assert result.metadata["status"] == "parser_unavailable", "旧格式语义一字未改"
        assert result.text == "", "旧格式不许产出任何正文（产出即冒充读到了）"
        note = file_reader.file_read_failure_note(result)
        assert note.startswith("[文件读取失败：") and "读不出来" in note
        assert phrase in note and "OLE2 复合文档" in note, "归因要点名它是哪一族"
        assert "本机" in note and "通路" in note, "要说本机有没有路（现算的）"
        assert "\\" not in note, "技术事实行不许出现路径"


def test_legacy_note_names_computed_missing_pieces(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """缺件代号是**现算**的：探测全打成空，措辞必须跟着变（写死的句子拦不住这一手）。"""
    target = tmp_path / "old.xls"
    target.write_bytes(_OLE2)
    live = file_reader.file_read_failure_note(file_reader.read_supported_file(target))
    _poison_probes(monkeypatch)
    blind = file_reader.file_read_failure_note(file_reader.read_supported_file(target))
    assert blind != live, "两段一模一样＝那句是写死的，不是探测出来的"
    assert "本机无转换通路" in blind
    assert "OfficeCOM(无)" in blind and "LibreOffice(无)" in blind
    assert "旧表读取器(无)" in blind


def test_misnamed_ppt_is_actually_read_by_magic(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """后缀 .ppt 内容 .pptx ⇒ 按魔数改走现代腿**真读**（从前一律被判「旧格式无适配器」）。"""
    pytest.importorskip("pptx")
    from pptx import Presentation

    real = tmp_path / "deck.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "守岸人读到了这页"
    deck.save(real)
    misnamed = tmp_path / "misnamed.ppt"
    misnamed.write_bytes(real.read_bytes())

    assert file_reader.sniff_container(misnamed) == "ooxml"
    result = file_reader.read_supported_file(misnamed)
    assert result.kind == "presentation"
    assert "守岸人读到了这页" in result.text, "能读的件不许被判读不了"
    assert not (result.metadata or {}).get("status")

    # 注毒：把嗅探改口成 OLE2，同一枚文件立刻回落到「无通路」诚实态——
    # 证明这条腿听的是魔数，不是后缀。
    monkeypatch.setattr(file_reader, "sniff_container", lambda _p: "ole2")
    poisoned = file_reader.read_supported_file(misnamed)
    assert poisoned.text == ""
    assert poisoned.metadata and poisoned.metadata["status"] == "parser_unavailable"


def test_misnamed_xls_reports_content_not_a_fake_ole2(tmp_path: Path) -> None:
    """.xls 内容实为 xlsx：本机 openpyxl 按后缀拒收 ⇒ 归因要说「内容是 OOXML、改道读过没读出来」，
    不许再谎称它是旧 OLE2，也不许静默交空正文。"""
    openpyxl = pytest.importorskip("openpyxl")
    real = tmp_path / "book.xlsx"
    book = openpyxl.Workbook()
    book.active["A1"] = "守岸人"
    book.save(real)
    misnamed = tmp_path / "misnamed.xls"
    misnamed.write_bytes(real.read_bytes())

    result = file_reader.read_supported_file(misnamed)
    metadata = result.metadata or {}
    fmt = str(metadata.get("format") or "")
    assert file_reader.sniff_container(misnamed) == "ooxml"
    assert "OLE2" not in fmt or "既非 OLE2" in fmt, "内容是 zip，就不许说它是 OLE2"
    assert "内容实为 OOXML 容器" in fmt and "改道读过" in fmt
    note = file_reader.file_read_failure_note(result)
    assert note, "读不出来必须开口，不许把空正文交给会话面当「文件是空的」"
    assert "读不出来" in note
    del metadata


def test_unrecognized_container_is_not_called_ole2(tmp_path: Path) -> None:
    """既非 OLE2 也非 OOXML：仍诚实降级，但不把没认出的容器说成旧格式。"""
    target = tmp_path / "junk.ppt"
    target.write_bytes(b"random bytes that are no known container" * 8)
    result = file_reader.read_supported_file(target)
    assert result.metadata and result.metadata["status"] == "parser_unavailable"
    fmt = str(result.metadata["format"])
    assert "既非 OLE2 也非 OOXML" in fmt and "不猜内容" in fmt
    assert result.text == ""


# --------------------------------------------------------------------------- ②


def test_new_legs_degrade_silently_on_a_blind_machine(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """缺能力机器（无 PATH 工具、无相关模块、无 COM）：新腿一律不接管、不抛、不装。"""
    _poison_probes(monkeypatch)
    assert file_reader.locate_text_layer_tool() == ""
    assert file_reader.second_pass_pdf_text(tmp_path / "any.pdf", 1000) == ""
    report = file_reader.local_capability_report()
    assert report["legacy_office_path"] is False
    assert report["scan_ocr_path"] is False
    assert report["text_layer_tools"] == []
    assert file_reader.scan_ocr_capability_code().endswith("装配点(空)")

    legacy = tmp_path / "old.xls"
    legacy.write_bytes(_OLE2)
    result = file_reader.read_supported_file(legacy)  # 不许抛
    assert result.metadata and result.metadata["status"] == "parser_unavailable"
    assert "本机无转换通路" in str(result.metadata["format"])


def test_scan_pdf_without_ocr_still_says_cannot_read(tmp_path: Path) -> None:
    """无文本层 PDF（本机无 OCR 组件、装配点空）⇒ 走既有四态，不许被当成「文件是空的」。"""
    pytest.importorskip("PIL")
    from PIL import Image

    target = tmp_path / "scan.pdf"
    Image.new("RGB", (240, 160), "white").save(target, format="PDF")

    result = file_reader.read_supported_file(target)
    scan = (result.metadata or {}).get("pdf_scan")
    if not isinstance(scan, dict) or scan.get("empty_reason") != "scanned_no_text":
        pytest.skip("本机 PIL/pypdf 没造出图像型无文字页，这一格由 test_pdf_scan_honesty 守")
    note = file_reader.file_read_failure_note(result)
    assert note.startswith("[文件无文字：")
    assert "OCR 兜底今天不生效" in note and "不是「没有内容」" in note
    assert "ocr通路=" in note, "缺件与装配点状态必须现算随行"
    assert "\\" not in note, "技术事实行不许出现路径"
    assert scan["ocr_path"], "代号不许是空话"


def test_scan_ocr_assembly_point_wiring_and_fault_containment(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """装配点（治「无装配点」那一格）：接线才动手，回字／抛错两形都在**真读取链**上验。"""
    pytest.importorskip("PIL")
    from PIL import Image

    scan_pdf = tmp_path / "scan.pdf"
    Image.new("RGB", (240, 160), "white").save(scan_pdf, format="PDF")
    baseline = file_reader.read_supported_file(scan_pdf)
    if ((baseline.metadata or {}).get("pdf_scan") or {}).get("empty_reason") != (
        "scanned_no_text"
    ):
        pytest.skip("本机 PIL/pypdf 没造出图像型无文字页")

    calls: list[int] = []

    def _handler(blobs: list[bytes]) -> str:
        calls.append(len(blobs))
        return "图像里识别出的字"

    def _boom(blobs: list[bytes]) -> str:
        calls.append(len(blobs))
        raise RuntimeError("外部 OCR 腿自己出的错")

    monkeypatch.setattr(
        file_reader, "_scan_page_image_blobs", lambda _r, _p: ([b"\xff\xd8x"], 1)
    )
    assert file_reader.scan_ocr_handler() is None, "缺省不许接管（接线前行为一字不变）"

    file_reader.set_scan_ocr_handler(_handler)
    try:
        code = file_reader.scan_ocr_capability_code()
        wired = file_reader.read_supported_file(scan_pdf)
    finally:
        file_reader.set_scan_ocr_handler(None)
    assert "装配点(已接线)" in code
    assert calls == [1], "接了线的钩子必须真被调用"
    assert "图像里识别出的字" in wired.text and "经装配点 OCR 读出" in wired.text
    wired_scan = (wired.metadata or {})["pdf_scan"]
    assert "empty_reason" not in wired_scan, "读回字了就不许再说这页没字"
    assert file_reader.file_read_failure_note(wired) == "", "读到了东西不许再说读不了"
    assert file_reader.scan_ocr_handler() is None, "卸下钩子必须回到缺省诚实行为"

    calls.clear()
    file_reader.set_scan_ocr_handler(_boom)
    try:
        broken = file_reader.read_supported_file(scan_pdf)
    finally:
        file_reader.set_scan_ocr_handler(None)
    broken_scan = (broken.metadata or {})["pdf_scan"]
    assert calls == [1]
    assert broken_scan["ocr_error"] == "RuntimeError", "钩子的错要记类名，不许撞穿入站"
    assert broken_scan["empty_reason"] == "scanned_no_text", "钩子失败要退回原诚实态"
    assert broken_scan["ocr_path"]
    note = file_reader.file_read_failure_note(broken)
    assert "OCR 兜底今天不生效" in note and "ocr通路=" in note


def test_second_pass_tool_budget_and_absence(tmp_path: Path) -> None:
    """pdftotext 那条腿：预算 0 不起进程；不是真 PDF 时回空串（空＝没读到，不是没有字）。"""
    ghost = tmp_path / "ghost.pdf"
    ghost.write_bytes(b"%PDF-1.4 not a real catalog\n" + b"0" * 200)
    assert file_reader.second_pass_pdf_text(ghost, 0) == ""


# --------------------------------------------------------------------------- ③


def test_text_whitelist_and_char_cap_untouched(tmp_path: Path) -> None:
    """文本白名单与 120000 字符上限＝本席只加识别，两把尺零回归（计数以真身为准）。"""
    for ext in (".xls", ".xlsx", ".ppt", ".pptx", ".doc", ".docx", ".pdf"):
        assert ext not in file_reader._TEXT_EXTS, f"{ext} 不许进文本白名单"
    for ext in (".tex", ".md", ".py", ".html", ".csv"):
        assert ext in file_reader._TEXT_EXTS, ext

    big = tmp_path / "big.txt"
    big.write_text("岸" * 300000, encoding="utf-8")
    body = file_reader.read_supported_file(big).text
    assert "只读取了文件开头约 120000 字符" in body, "截断必须显式说明"
    assert "没读不等于没有内容" in body
    assert len(body.split("\n[")[0]) == 120000, "字符上限不许被改动"
