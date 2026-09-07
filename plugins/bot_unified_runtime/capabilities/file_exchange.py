"""文件收发能力（接收代码/Markdown 调试 + 导出多格式文档）。

接收：管理员上传 .py 等文件 → 语法检查 + 受限子进程运行 → 回报错误与输出。
子进程无 shell、隔离模式（-I）、限时、工作目录为临时目录、不继承用户环境。

发送：``/bot 文件 <md|docx|pptx|xlsx|pdf> <主题>`` → LLM 生成 Markdown →
转换为目标格式 → 以平台上传文件接口发送。文档库缺失时诚实说明，不伪造。
"""

from __future__ import annotations

import os
import re
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

EXPORT_FORMATS = ("md", "docx", "pptx", "xlsx", "pdf")
_RUNNABLE_EXTENSIONS = {".py"}
_TEXT_EXTENSIONS = {".md", ".markdown", ".txt", ".json", ".yaml", ".yml", ".toml", ".csv"}
_FILE_EXPORT_RE = re.compile(
    r"^(?:/bot\s+)?文件\s+(?P<fmt>md|markdown|docx|pptx|xlsx|pdf)\s+(?P<topic>\S.*)$",
    re.IGNORECASE,
)
_DOCUMENT_PROMPT = (
    "你是文档撰写助手。围绕用户给出的主题撰写一份结构清晰的中文 Markdown 文档："
    "用 #/## 分级标题组织，适量使用 - 列表和 | 表格 |，正文 600-1200 字，"
    "不要输出代码块围栏以外的内容，直接输出 Markdown 本身。"
)


def is_file_export_command(text: str) -> bool:
    return _FILE_EXPORT_RE.match(text.strip()) is not None


def parse_file_export_command(text: str) -> tuple[str, str] | None:
    match = _FILE_EXPORT_RE.match(text.strip())
    if match is None:
        return None
    fmt = match.group("fmt").lower()
    return ("md", match.group("topic").strip()) if fmt == "markdown" else (
        fmt,
        match.group("topic").strip(),
    )


def run_code_debug(
    file_path: str | Path,
    *,
    timeout_seconds: float = 15.0,
    python_executable: str = "",
) -> str:
    """受限运行代码文件并返回调试报告（只含退出码与 stdout/stderr 摘要）。"""
    path = Path(file_path)
    if not path.exists():
        return "文件不存在或无法读取。"
    suffix = path.suffix.lower()
    if suffix not in _RUNNABLE_EXTENSIONS:
        if suffix not in _TEXT_EXTENSIONS:
            return f"已接收 {path.name}。该类型暂不支持调试，仅做接收确认。"
        text = path.read_text(encoding="utf-8", errors="replace")
        return (
            f"已接收 {path.name}（{len(text)} 字符 / {len(text.splitlines())} 行）。"
            "该类型暂不支持直接运行。"
        )
    source = path.read_text(encoding="utf-8", errors="replace")
    try:
        compile(source, path.name, "exec")
    except SyntaxError as exc:
        return f"语法错误：第 {exc.lineno} 行：{exc.msg}"
    executable = python_executable or sys.executable
    env = {
        "PYTHONDONTWRITEBYTECODE": "1",
        "PATH": os.environ.get("PATH", ""),
        "SYSTEMROOT": os.environ.get("SYSTEMROOT", ""),
        "COMSPEC": os.environ.get("COMSPEC", ""),
    }
    try:
        with tempfile.TemporaryDirectory(prefix="bot_debug_") as workdir:
            try:
                completed = subprocess.run(  # noqa: PLW1510 - 非零退出码是调试结果，不抛异常。
                    [executable, "-I", str(path.resolve())],
                    cwd=workdir,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=timeout_seconds,
                    env=env,
                )
            except subprocess.TimeoutExpired:
                return f"语法检查通过 ✓；运行超时（>{timeout_seconds:.0f}s），已终止。"
            except OSError as exc:
                return f"语法检查通过 ✓；运行失败：{type(exc).__name__}"
    except OSError as exc:
        return f"运行环境异常：{type(exc).__name__}"
    parts = ["语法检查通过 ✓"]
    if completed.returncode == 0:
        parts.append("运行成功（退出码 0）")
    else:
        parts.append(f"运行失败（退出码 {completed.returncode}）")
    stdout_tail = (completed.stdout or "").strip()[-400:]
    stderr_tail = (completed.stderr or "").strip()[-600:]
    if stdout_tail:
        parts.append(f"stdout 末尾：\n{stdout_tail}")
    if stderr_tail:
        parts.append(f"stderr（错误定位）：\n{stderr_tail}")
    return "\n".join(parts)


@dataclass
class DocumentBlock:
    """Markdown 块级元素（导出器的中间表示）。"""

    kind: str  # heading / bullet / code / para / table
    level: int = 0
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)


_TABLE_ROW_RE = re.compile(r"^\s*\|(.+)\|\s*$")


def parse_markdown_blocks(markdown: str) -> list[DocumentBlock]:
    blocks: list[DocumentBlock] = []
    in_code = False
    code_lines: list[str] = []
    for raw_line in str(markdown or "").splitlines():
        line = raw_line.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                blocks.append(
                    DocumentBlock(kind="code", text="\n".join(code_lines))
                )
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line.strip())
        if heading:
            blocks.append(
                DocumentBlock(kind="heading", level=len(heading.group(1)), text=heading.group(2).strip())
            )
            continue
        table_row = _TABLE_ROW_RE.match(line)
        if table_row:
            cells = [cell.strip() for cell in table_row.group(1).split("|")]
            if blocks and blocks[-1].kind == "table" and set("".join(cells)) <= set("-: "):
                continue  # 分隔行
            if blocks and blocks[-1].kind == "table":
                blocks[-1].rows.append(cells)
            else:
                blocks.append(DocumentBlock(kind="table", rows=[cells]))
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line.strip())
        if bullet:
            blocks.append(DocumentBlock(kind="bullet", text=bullet.group(1).strip()))
            continue
        blocks.append(DocumentBlock(kind="para", text=line.strip()))
    if in_code and code_lines:
        blocks.append(DocumentBlock(kind="code", text="\n".join(code_lines)))
    return blocks


def _safe_file_name(title: str, fmt: str) -> str:
    slug = re.sub(r"[^\w\u4e00-\u9fff-]+", "_", str(title or "document").strip())[:40]
    slug = slug.strip("_") or "document"
    return f"{slug}.{fmt}"


def _find_cjk_font() -> str:
    candidates = (
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simsun.ttc",
    )
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    return ""


def export_document(markdown: str, fmt: str, out_dir: Path, *, title: str = "") -> tuple[Path, str]:
    """Markdown → 目标格式；返回 (文件路径, 错误说明)。错误非空即失败。"""
    fmt = str(fmt or "").lower()
    if fmt == "markdown":
        fmt = "md"
    if fmt not in EXPORT_FORMATS:
        return Path(), f"不支持的格式：{fmt}"
    out_dir = Path(out_dir)
    try:
        out_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        return Path(), f"输出目录不可写：{type(exc).__name__}"
    target = out_dir / _safe_file_name(title, fmt)
    blocks = parse_markdown_blocks(markdown)
    try:
        if fmt == "md":
            target.write_text(str(markdown or ""), encoding="utf-8")
            return target, ""
        if fmt == "docx":
            return _export_docx(blocks, target, title)
        if fmt == "xlsx":
            return _export_xlsx(blocks, target)
        if fmt == "pptx":
            return _export_pptx(blocks, target, title)
        if fmt == "pdf":
            return _export_pdf(blocks, target, title)
    except ImportError as exc:
        return Path(), f"缺少文档库（{exc.name or '依赖'}），无法生成 {fmt.upper()}"
    except Exception as exc:  # noqa: BLE001 - 转换失败只回报类型与库级摘要，不回显内容。
        return Path(), f"转换失败：{type(exc).__name__}: {str(exc)[:120]}"
    return Path(), f"不支持的格式：{fmt}"


def _export_docx(blocks: list[DocumentBlock], target: Path, title: str) -> tuple[Path, str]:
    import docx  # python-docx

    document = docx.Document()
    if title:
        document.add_heading(title, level=0)
    for block in blocks:
        if block.kind == "heading":
            document.add_heading(block.text, level=max(1, min(4, block.level)))
        elif block.kind == "bullet":
            document.add_paragraph(block.text, style="List Bullet")
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            run.font.name = "Courier New"
        elif block.kind == "table":
            table = document.add_table(rows=len(block.rows), cols=max(len(r) for r in block.rows))
            for row_index, row in enumerate(block.rows):
                for col_index, cell in enumerate(row):
                    table.rows[row_index].cells[col_index].text = cell
        else:
            document.add_paragraph(block.text)
    document.save(str(target))
    return target, ""


def _export_xlsx(blocks: list[DocumentBlock], target: Path) -> tuple[Path, str]:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    row_index = 1
    wrote_table = False
    for block in blocks:
        if block.kind == "table":
            for row in block.rows:
                for col_index, cell in enumerate(row, start=1):
                    sheet.cell(row=row_index, column=col_index, value=cell)
                row_index += 1
            wrote_table = True
            row_index += 1
        else:
            sheet.cell(row=row_index, column=1, value=block.text)
            row_index += 1
    if not wrote_table:
        sheet.title = "内容"
    workbook.save(str(target))
    return target, ""


def _export_pptx(blocks: list[DocumentBlock], target: Path, title: str) -> tuple[Path, str]:
    from pptx import Presentation
    from pptx.util import Pt

    presentation = Presentation()
    if title:
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = title
    current_bullets: list[str] = []
    current_heading = ""

    def _flush() -> None:
        nonlocal current_bullets, current_heading
        if not current_bullets:
            return
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = current_heading or title or "内容"
        body = slide.placeholders[1].text_frame
        for index, item in enumerate(current_bullets[:8]):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = item
            paragraph.font.size = Pt(18)
        current_bullets = []

    for block in blocks:
        if block.kind == "heading":
            _flush()
            current_heading = block.text
        elif block.kind == "bullet" or (
            block.kind == "para" and len(current_bullets) < 8
        ):
            current_bullets.append(block.text)
    _flush()
    if not presentation.slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = title or "空文档"
    presentation.save(str(target))
    return target, ""


def _export_pdf(blocks: list[DocumentBlock], target: Path, title: str) -> tuple[Path, str]:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    font_path = _find_cjk_font()
    if font_path:
        pdf.add_font("cjk", "", font_path)
        pdf.set_font("cjk", size=12)
    else:
        # 无中文字体时降级：仅保证不崩溃，报告由调用方提示。
        pdf.set_font("helvetica", size=12)

    def _line(height: float, text: str) -> None:
        # fpdf2 的 multi_cell 默认 new_x=RIGHT 不回左边距，必须手动复位，
        # 否则连续两行会因剩余宽度为 0 抛“无横向空间”。
        pdf.multi_cell(0, height, text)
        pdf.set_x(pdf.l_margin)

    if title:
        pdf.set_font_size(18)
        _line(10, title)
        pdf.set_font_size(12)
        pdf.ln(2)
    for block in blocks:
        if block.kind == "heading":
            pdf.set_font_size(12 + max(0, 5 - block.level) * 2)
            _line(8, block.text)
            pdf.set_font_size(12)
        elif block.kind == "bullet":
            # 用 ASCII 项目符号：中文字体缺 U+2022 字形时 fpdf2 会按零宽度抛错。
            _line(7, f"- {block.text}")
        elif block.kind == "code":
            for line in block.text.splitlines():
                _line(6, f"    {line}")
        elif block.kind == "table":
            for row in block.rows:
                _line(7, "  ".join(row))
        else:
            _line(7, block.text)
        pdf.ln(1)
    pdf.output(str(target))
    return target, ""
