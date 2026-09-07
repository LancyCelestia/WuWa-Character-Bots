"""Safe, non-executing readers and generated-file helpers."""
from __future__ import annotations

import ast
import re
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

_TEXT_EXTS={".txt",".log",".md",".csv",".json",".yaml",".yml",".py",".c",".h",".cpp",".hpp",".cc",".cxx",".cs",".java",".js",".ts",".tsx",".jsx",".go",".rs",".php",".sh",".ps1",".sql",".html",".css",".xml"}
_CODE_EXTS={".py",".c",".h",".cpp",".hpp",".cc",".cxx",".cs",".java",".js",".ts",".tsx",".jsx",".go",".rs",".php",".sh",".ps1",".sql",".html",".css"}
@dataclass(frozen=True)
class FileReadResult:
    path: Path
    kind: str
    text: str
    title: str = ""
    metadata: dict[str, Any] | None = None
@dataclass(frozen=True)
class GeneratedFile:
    path: Path
    kind: str

def _text(path: Path, max_chars: int) -> str:
    raw=path.read_bytes()[:max_chars*4]
    if b"\x00" in raw[:4096]: return ""
    return raw.decode("utf-8",errors="replace")[:max_chars]

def read_supported_file(path: str | Path, *, max_chars: int = 120000) -> FileReadResult:
    source=Path(path).expanduser()
    if not source.is_file(): return FileReadResult(source,"missing","")
    ext=source.suffix.lower()
    if ext in _TEXT_EXTS:
        return FileReadResult(source,"code" if ext in _CODE_EXTS else "text",_text(source,max_chars),source.name)
    if ext == ".docx":
        try:
            from docx import Document
            text="\n".join(p.text for p in Document(str(source)).paragraphs)
            return FileReadResult(source,"document",text[:max_chars],source.name)
        except (ImportError, OSError, ValueError, TypeError): return FileReadResult(source,"document","")
    if ext in {".xlsx",".xls"}:
        try:
            import openpyxl
            book=openpyxl.load_workbook(source,read_only=True,data_only=True)
            rows=[]
            for sheet in book.worksheets:
                rows.append(f"[工作表：{sheet.title}]")
                for row in sheet.iter_rows(values_only=True): rows.append(" | ".join("" if v is None else str(v) for v in row))
            return FileReadResult(source,"spreadsheet","\n".join(rows)[:max_chars],source.name)
        except (ImportError, OSError, ValueError, TypeError): return FileReadResult(source,"spreadsheet","")
    if ext in {".pptx",".ppt"}:
        try:
            from pptx import Presentation
            lines=[]
            for index, slide in enumerate(Presentation(str(source)).slides,1):
                lines.append(f"[幻灯片 {index}]")
                lines.extend(shape.text for shape in slide.shapes if hasattr(shape,"text"))
            return FileReadResult(source,"presentation","\n".join(lines)[:max_chars],source.name)
        except (ImportError, OSError, ValueError, TypeError): return FileReadResult(source,"presentation","")
    return FileReadResult(source,"unknown","")

def artifact_request(user_text: str) -> tuple[str, str] | None:
    text = (user_text or '').lower()
    if not re.search(r'生成|写入|保存|导出|创建|制作|generate|write|save|export|create', text):
        return None
    languages = {'python':'py', 'c++':'cpp', 'c#':'cs', 'java':'java', 'javascript':'js',
                 'typescript':'ts', 'powershell':'ps1', 'bash':'sh'}
    if re.search(r'代码|脚本|code|script', text) or any(k in text for k in languages):
        ext = next((v for k,v in sorted(languages.items(),key=lambda kv:-len(kv[0])) if k in text), 'txt')
        return 'code', ext
    if re.search(r'文件|文档|txt|文本|document|file|提示词|人设',text):
        ext = 'txt' if 'txt' in text or '文本' in text else 'md'
        if re.search(r'\b(docx|xlsx|pptx|pdf)\b|word|excel|powerpoint', text):
            return 'unsupported', 'txt'
        return 'document', ext
    return None


def build_generated_file(user_text: str, reply_text: str, output_dir: str | Path) -> GeneratedFile | None:
    """Consume raw model output, not presentation text. Write once with unique names."""
    intent = artifact_request(user_text)
    if intent is None:
        return None
    kind, ext = intent
    if kind == 'unsupported':
        raise ValueError('此入口目前支持 TXT、Markdown 和代码文件；不能把纯文本冒充 Office/PDF。')
    blocks = list(re.finditer(r'```([A-Za-z0-9+#-]*)[^\S\n]*\n(.*?)```',reply_text or '',re.DOTALL))
    if kind == 'code':
        if not blocks:
            # Bare Python is accepted only if syntactically complete; never strip quotes.
            if ext != 'py':
                raise ValueError('模型没有提供完整代码块，本次没有生成文件。')
            content = reply_text
        elif len(blocks) != 1:
            raise ValueError('模型返回多个代码块，请要求生成单个完整文件。')
        else:
            content = blocks[0][2]
            if ext == 'txt':
                ext = {'python':'py','cpp':'cpp','c++':'cpp','c#':'cs','javascript':'js','java':'java'}.get(blocks[0][1].lower(),'txt')
        if ext == 'py':
            try:
                ast.parse(content)
            except SyntaxError as exc:
                raise ValueError('模型代码未通过 Python 语法检查，本次没有生成文件。') from exc
    else:
        content = blocks[0][2] if len(blocks) == 1 else reply_text
    if not content.strip():
        raise ValueError('模型返回空文件内容。')
    if len(content.encode('utf-8')) > 2 * 1024 * 1024:
        raise ValueError('生成内容超过 2 MiB 文件限制。')
    directory = Path(output_dir).resolve()
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f'generated_{kind}_{uuid.uuid4().hex[:12]}.{ext}'
    with path.open('x', encoding='utf-8', newline='') as stream:
        stream.write(content.rstrip('\r\n') + '\n')
    return GeneratedFile(path, kind)
