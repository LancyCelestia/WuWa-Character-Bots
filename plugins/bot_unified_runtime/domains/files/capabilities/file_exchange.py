"""文件收发能力（接收代码/Markdown 的读码反馈 + 导出多格式文档）。

接收：管理员上传 .py 等文件 → **语法检查 + 静态结构摘要**（顶层定义、导入、语句量）
→ 回报理解。默认**不执行**（``_CODE_EXECUTION_ENABLED=False``）：子进程不是沙箱——
``-I`` 只隔离 Python 环境（不继承用户 site/环境变量），无 shell、限时、工作目录为临时目录，
但进程仍以 bot 账户的全部 OS 权限运行，可访问 bot 能访问的任意文件与网络。
执行能力要恢复，必须等统一安全与执行引擎（Job Object 配额 + 超管书面同意）就位，
详见 ``docs/design/safety-execution-engine-spec.md`` §6 与 Wave 2。

发送：``/bot 文件 <md|docx|pptx|xlsx|pdf> <主题>`` → LLM 生成 Markdown →
转换为目标格式 → 以平台上传文件接口发送。文档库缺失时诚实说明，不伪造。

**落盘收口（需求 16(2)，2026-09-25 T-FILE-2）**：本件不再自己拼目标路径写盘。
「创建文件 / 修改文件」两个动词的唯一落盘口是
``domains/files/sender/restricted_runner.py``（白名单落点 + 单文件与每日限额 +
扩展名与字节指纹双查 + 段名消毒 + ``resolve()`` 后越界复核 + 原子发布）。
转换库今天只往运行器**签发的暂存位**写字节，验收不过就碰不到目的地。

**创建与执行是两条开关**：``_CODE_EXECUTION_ENABLED``（见下方注释）关的是「跑上传
的代码」，与「写一个 ``.py`` 落盘」无关——关态下创建照常成功、执行次数恒为 0，这条
由 ``tests/test_file_exchange_restricted_runner.py`` 的「创建但不执行」锁钉住，不靠
注释与自觉。
"""

from __future__ import annotations

import ast
import hashlib
import logging
import os
import re
import subprocess
import sys
import tempfile
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from dataclasses import replace as dataclass_replace
from pathlib import Path
from typing import Any

from plugins.bot_unified_runtime.domains.chat_reply.capabilities import user_copy
from plugins.bot_unified_runtime.domains.core.safety_exec import policy as safety_policy
from plugins.bot_unified_runtime.domains.core.safety_exec.action_catalog import ActionId
from plugins.bot_unified_runtime.domains.files.sender.restricted_runner import (
    MODE_CREATE_OR_REPLACE,
    DailyQuotaLedger,
    WriteLimits,
    WriteOutcome,
    WritePolicy,
    create_bytes,
    plain_reason,
    policy_for_roots,
    read_confined_bytes,
    replace_bytes,
    resolve_existing,
    revise_in_place,
    sendable_verdict,
    stage_write,
)

EXPORT_FORMATS = ("md", "docx", "pptx", "xlsx", "pdf")

logger = logging.getLogger(__name__)
_RUNNABLE_EXTENSIONS = {".py"}
_TEXT_EXTENSIONS = {".md", ".markdown", ".txt", ".json", ".yaml", ".yml", ".toml", ".csv"}
_FILE_EXPORT_RE = re.compile(
    r"^(?:/bot\s+)?文件\s+(?P<fmt>md|markdown|docx|pptx|xlsx|pdf)\s+(?P<topic>\S.*)$",
    re.IGNORECASE,
)
_DOCUMENT_PROMPT = (
    "你是文档撰写助手。围绕用户给出的主题撰写一份结构清晰的中文 Markdown 文档："
    "用 #/## 分级标题组织，适量使用 - 列表和 | 表格 |，正文 600-1200 字，"
    "不要输出代码块围栏以外的内容，直接输出 Markdown 本身。"
)


def is_file_export_command(text: str) -> bool:
    return _FILE_EXPORT_RE.match(text.strip()) is not None


def parse_file_export_command(text: str) -> tuple[str, str] | None:
    match = _FILE_EXPORT_RE.match(text.strip())
    if match is None:
        return None
    fmt = match.group("fmt").lower()
    return ("md", match.group("topic").strip()) if fmt == "markdown" else (
        fmt,
        match.group("topic").strip(),
    )


#: 执行开关（2026-09-25 SAFE-EXEC Wave 1·P-1 甲止血）。
#: 真身理由：此前任何在 ``bot_admin_user_ids`` 里的账号，往 bot 所在群或私聊丢一个 .py
#: 就以 **bot 进程的 OS 权限**直接执行（``-I`` 只隔离 Python 环境，不隔离文件/网络/进程），
#: 一个管理员号被盗＝本机被接管。统一安全与执行引擎（Wave 2：Job Object 配额 +
#: 超管书面同意）落地之前，这一路只做工号检查（语法 + 静态结构摘要），不再起跑。
#: 刻意不做成配置键：装配期快照式的开关看着能热改、实则不会生效，且它属于安全红线一侧。
_CODE_EXECUTION_ENABLED = False


def _static_structure_report(source: str, name: str) -> str:
    """静态读码报告：顶层定义、导入、语句量——不跑一行。"""
    try:
        tree = ast.parse(source, filename=name)
    except SyntaxError as exc:  # 上游已拦过一次，这里只兜底
        return f"语法错误：第 {exc.lineno} 行：{exc.msg}"
    imports: list[str] = []
    defs: list[str] = []
    for node in tree.body:
        if isinstance(node, (ast.Import, ast.ImportFrom)):
            module = getattr(node, "module", "") or ""
            names = ", ".join(alias.name for alias in node.names)
            imports.append(f"{module or names}")
        elif isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            kind = "类" if isinstance(node, ast.ClassDef) else "函数"
            defs.append(f"{kind} {node.name}")
    parts = [
        "语法检查通过 ✓",
        f"结构摘要：{len(tree.body)} 个顶层语句；"
        + (f"定义 {len(defs)} 个（{'; '.join(defs[:6])}{'…' if len(defs) > 6 else ''}）" if defs else "无顶层定义"),
    ]
    if imports:
        parts.append("导入：" + "、".join(imports[:10]) + ("…" if len(imports) > 10 else ""))
    parts.append(
        "本轮未执行：跑代码要先过统一执行引擎（资源配额 + 主人书面同意），"
        "现在只做工号与结构检查。要我继续往下读逻辑或找问题，直接说。"
    )
    return "\n".join(parts)


def _confined_debug_read_text(
    path: Path, policy: WritePolicy | None
) -> str | WriteOutcome:
    """代码检查腿的受限回读：走唯一回读口 ``read_confined_bytes``，不自己开文件。

    口径（S-16-READBACK 的修法，2026-09-26 S-FILES-LAND）：缺省策略取**该文件所在
    目录为单根**——这份字节本来已由入站腿落在那儿，本函数只负责读，不负责造第二
    份 containment 判据；白名单 containment、禁触名册直判、可执行形态与单文件限额
    全部照问咽喉（旧写法是 ``Path.read_text`` 直读整份文件、零上限，500MB 的 .py
    会被整份吃进内存再 ``ast.parse``）。装配层（施工单落根文件后）把
    ``write_policy_from_config(config)`` 交进来时改用调用方策略，根不随文件走。
    """
    active = policy or policy_for_roots([path.parent])
    read_back = read_confined_bytes(path.name, policy=active)
    if isinstance(read_back, WriteOutcome):
        return read_back
    data, _digest = read_back
    return data.decode("utf-8", errors="replace")


def _confined_read_failure_line(name: str, failure: WriteOutcome) -> str:
    """受限回读被拒时的人话（只点名文件名词与拒绝代号，不带路径明文）。"""
    return (
        f"已接收 {name}，但这份文件我按受限回读口径读不了："
        f"{failure.error_message() or plain_reason(failure.reason_code)}"
    )


def run_code_debug(
    file_path: str | Path,
    *,
    timeout_seconds: float = 15.0,
    python_executable: str = "",
    execute: bool | None = None,
    policy: WritePolicy | None = None,
) -> str:
    """检查代码文件并返回报告；**缺省不执行**（见 ``_CODE_EXECUTION_ENABLED``）。

    非沙箱：``execute=True`` 那条路仍以 bot 进程同等 OS 权限运行（``-I`` 只隔离
    Python 环境），所以它只应被统一安全与执行引擎在拿到超管书面同意后调用，
    生产装配路径（根文件）不传这个参数。
    """
    path = Path(file_path)
    if not path.exists():
        return "文件不存在或无法读取。"
    suffix = path.suffix.lower()
    if suffix not in _RUNNABLE_EXTENSIONS and suffix not in _TEXT_EXTENSIONS:
        return f"已接收 {path.name}。该类型暂不支持调试，仅做接收确认。"
    payload = _confined_debug_read_text(path, policy)
    if isinstance(payload, WriteOutcome):
        return _confined_read_failure_line(path.name, payload)
    text = payload
    if suffix not in _RUNNABLE_EXTENSIONS:
        return (
            f"已接收 {path.name}（{len(text)} 字符 / {len(text.splitlines())} 行）。"
            "该类型暂不支持直接运行。"
        )
    source = text
    try:
        compile(source, path.name, "exec")
    except SyntaxError as exc:
        return f"语法错误：第 {exc.lineno} 行：{exc.msg}"
    should_execute = _CODE_EXECUTION_ENABLED if execute is None else bool(execute)
    if not should_execute:
        return _static_structure_report(source, path.name)
    executable = python_executable or sys.executable
    env = {
        "PYTHONDONTWRITEBYTECODE": "1",
        "PATH": os.environ.get("PATH", ""),
        "SYSTEMROOT": os.environ.get("SYSTEMROOT", ""),
        "COMSPEC": os.environ.get("COMSPEC", ""),
    }
    try:
        # 审计#19：Windows 下孙进程可能仍锁着 workdir，with 退出清理抛 OSError
        # 会经外层 except 把真实运行结果吞成"运行环境异常"；ignore_cleanup_errors
        # 让清理失败静默，结果照常交付。
        with tempfile.TemporaryDirectory(
            prefix="bot_debug_", ignore_cleanup_errors=True
        ) as workdir:
            try:
                completed = subprocess.run(  # noqa: PLW1510 - 非零退出码是调试结果，不抛异常。
                    [executable, "-I", str(path.resolve())],
                    cwd=workdir,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=timeout_seconds,
                    env=env,
                )
            except subprocess.TimeoutExpired:
                return f"语法检查通过 ✓；运行超时（>{timeout_seconds:.0f}s），已终止。"
            except OSError as exc:
                logger.warning("run_code_debug: failed to start python: %s", exc)
                return f"语法检查通过 ✓；本机没能启动 Python 跑这段代码，这次运行不了。{user_copy.RUN_ENV_FAILURE_ADVICE}"
    except OSError as exc:
        logger.warning("run_code_debug: failed to prepare workdir: %s", exc)
        return f"本机没法准备运行用的临时目录，这次跑不了。{user_copy.RUN_ENV_FAILURE_ADVICE}"
    parts = ["语法检查通过 ✓"]
    if completed.returncode == 0:
        parts.append("运行成功（退出码 0）")
    else:
        parts.append(f"运行失败（退出码 {completed.returncode}）")
    stdout_tail = (completed.stdout or "").strip()[-400:]
    stderr_tail = (completed.stderr or "").strip()[-600:]
    if stdout_tail:
        parts.append(f"stdout 末尾：\n{stdout_tail}")
    if stderr_tail:
        parts.append(f"stderr（错误定位）：\n{stderr_tail}")
    return "\n".join(parts)


@dataclass
class DocumentBlock:
    """Markdown 块级元素（导出器的中间表示）。"""

    kind: str  # heading / bullet / code / para / table
    level: int = 0
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)


_TABLE_ROW_RE = re.compile(r"^\s*\|(.+)\|\s*$")


def parse_markdown_blocks(markdown: str) -> list[DocumentBlock]:
    blocks: list[DocumentBlock] = []
    in_code = False
    code_lines: list[str] = []
    for raw_line in str(markdown or "").splitlines():
        line = raw_line.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                blocks.append(
                    DocumentBlock(kind="code", text="\n".join(code_lines))
                )
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line.strip())
        if heading:
            blocks.append(
                DocumentBlock(kind="heading", level=len(heading.group(1)), text=heading.group(2).strip())
            )
            continue
        table_row = _TABLE_ROW_RE.match(line)
        if table_row:
            cells = [cell.strip() for cell in table_row.group(1).split("|")]
            if blocks and blocks[-1].kind == "table" and set("".join(cells)) <= set("-: "):
                continue  # 分隔行
            if blocks and blocks[-1].kind == "table":
                blocks[-1].rows.append(cells)
            else:
                blocks.append(DocumentBlock(kind="table", rows=[cells]))
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line.strip())
        if bullet:
            blocks.append(DocumentBlock(kind="bullet", text=bullet.group(1).strip()))
            continue
        blocks.append(DocumentBlock(kind="para", text=line.strip()))
    if in_code and code_lines:
        blocks.append(DocumentBlock(kind="code", text="\n".join(code_lines)))
    return blocks


def _safe_file_name(title: str, fmt: str) -> str:
    normalized = str(title or "document").strip()
    slug = re.sub(r"[^\w\u4e00-\u9fff-]+", "_", normalized)[:40]
    slug = slug.strip("_") or "document"
    # 审计#36：截断归一化后不同主题可能得到同一 slug，静默覆盖旧导出；
    # 追加主题内容短 hash 区分。
    digest = hashlib.sha1(normalized.encode("utf-8")).hexdigest()[:8]
    return f"{slug}_{digest}.{fmt}"


def _find_cjk_font() -> str:
    candidates = (
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simsun.ttc",
    )
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    return ""


def export_document(
    markdown: str,
    fmt: str,
    out_dir: Path,
    *,
    title: str = "",
    policy: WritePolicy | None = None,
    ledger: DailyQuotaLedger | None = None,
    date_key: Callable[[], str] | None = None,
    staging_dir: Path | str | None = None,
) -> tuple[Path, str]:
    """Markdown → 目标格式 → **经受限运行器**落盘；返回 (文件路径, 错误说明)。

    错误说明非空即失败。本函数**不碰目的地**：转换库写进运行器签发的暂存位，
    验收与原子发布全在 ``restricted_runner`` 一处（需求 16(2)）。

    缺省策略＝「只许写进 ``out_dir`` 这一个根、限额按运行器缺省」。装配层要收窄
    扩展名、改限额或接禁触名册（``external_verdict``），显式传 ``policy``。
    同主题重导＝同名：走 ``create_or_replace``，动词由运行器判并回写结论
    （旧的「静默覆盖」语义保持，但不再是无人判的覆盖）。

    ``staging_dir`` 只给测试用（把暂存位收进 ``tmp_path``，别污染系统临时目录）。
    """
    fmt = str(fmt or "").lower()
    if fmt == "markdown":
        fmt = "md"
    if fmt not in EXPORT_FORMATS:
        return Path(), f"不支持的格式：{fmt}"
    active_policy = policy or policy_for_roots([out_dir])
    blocks = parse_markdown_blocks(markdown)
    staged = stage_write(_safe_file_name(title, fmt), staging_dir=staging_dir)
    if isinstance(staged, WriteOutcome):
        return Path(), staged.error_message()
    try:
        error = ""
        try:
            if fmt == "md":
                staged.accept_text(str(markdown or ""))
            elif fmt == "docx":
                error = _export_docx(blocks, staged.path, title)[1]
            elif fmt == "xlsx":
                error = _export_xlsx(blocks, staged.path)[1]
            elif fmt == "pptx":
                error = _export_pptx(blocks, staged.path, title)[1]
            else:
                error = _export_pdf(blocks, staged.path, title)[1]
        except ImportError as exc:
            return Path(), f"缺少文档库（{exc.name or '依赖'}），无法生成 {fmt.upper()}"
        except Exception as exc:  # noqa: BLE001 - 转换失败只回报类型与库级摘要，不回显内容。
            return Path(), f"转换失败：{type(exc).__name__}: {str(exc)[:120]}"
        if error:
            return Path(), error
        outcome = staged.publish(
            active_policy,
            mode=MODE_CREATE_OR_REPLACE,
            ledger=ledger,
            date_key=date_key,
        )
        if not outcome.ok:
            return Path(), outcome.error_message()
        return outcome.path or Path(), ""
    finally:
        staged.cleanup()


def create_document_file(
    name: str,
    content: str | bytes,
    *,
    policy: WritePolicy,
    ledger: DailyQuotaLedger | None = None,
    date_key: Callable[[], str] | None = None,
    staging_dir: Path | str | None = None,
) -> WriteOutcome:
    """「创建文件」的唯一能力层入口（绝不覆盖同名件）。

    只声明要写什么，**不拼路径、不开文件**：白名单、限额、扩展名与字节指纹、
    越界复核、配额全在运行器。拿不到白名单根时返回 ``no_whitelist``（fail-closed）。
    """
    return create_bytes(
        name,
        _as_bytes(content),
        policy=policy,
        ledger=ledger,
        date_key=date_key,
        staging_dir=staging_dir,
    )


def revise_document_file(
    name: str,
    content: str | bytes,
    *,
    policy: WritePolicy,
    ledger: DailyQuotaLedger | None = None,
    date_key: Callable[[], str] | None = None,
    staging_dir: Path | str | None = None,
) -> WriteOutcome:
    """「修改文件」的唯一能力层入口（目标不在即 ``target_missing``，绝不新建）。"""
    return replace_bytes(
        name,
        _as_bytes(content),
        policy=policy,
        ledger=ledger,
        date_key=date_key,
        staging_dir=staging_dir,
    )


def _as_bytes(content: str | bytes) -> bytes:
    return content if isinstance(content, bytes) else str(content or "").encode("utf-8")


# ---------------------------------------------------------------------------
# 写盘装配口（需求 16(2) 容器 + 16(4) 唯一裁决出口；2026-09-26 S-FILES-LAND）
#
# 六枚 ``bot_files_*`` 配置键的唯一生产读点就在本节；根 matcher（施工单在席位
# 日志 §施工单一）只做四件事：管理员门 → ``parse_file_revise_command`` → 按指令
# 产出 ``transform`` → 把 ``run_document_*`` 结果里的 ``file_parts`` 交统一管线。
# 本口不投递、不拼路径、不复制判据：「这一步谁够格做」在 ``safety_exec.policy.
# decide`` 问一次（文件域唯一裁决出口），「落在哪、能不能落」在
# ``restricted_runner`` 问——两道闸都不长第三家。
# ---------------------------------------------------------------------------

_FILE_REVISE_RE = re.compile(
    r"^(?:/bot\s+)?改文件\s+(?P<name>\S+)\s+(?P<instruction>\S.*)$",
    re.IGNORECASE,
)

_WRITE_DISABLED_REPLY = "文件写盘口今天没开（BOT_FILES_WRITE_ENABLED=false），这次不落盘。"


def is_file_revise_command(text: str) -> bool:
    return _FILE_REVISE_RE.match(str(text or "").strip()) is not None


def parse_file_revise_command(text: str) -> tuple[str, str] | None:
    """``改文件 <名> <指令>`` → ``(名, 指令)``；不认返回 ``None``。

    名只当**相对落点**交给咽喉消毒定位（绝不原样拼路径），指令整段交给装配层
    构造 transform——本函数不猜文件名、不做存在性回答（oracle 纪律同
    ``resolve_existing``）。
    """
    match = _FILE_REVISE_RE.match(str(text or "").strip())
    if match is None:
        return None
    return match.group("name").strip(), match.group("instruction").strip()


@dataclass(frozen=True)
class DocumentOpResult:
    """写盘装配口的会话面结论：人话 + 可选出站件；**不含盘符路径明文**。"""

    ok: bool
    reply_text: str
    reason_code: str = ""
    file_parts: list[dict[str, str]] = field(default_factory=list)


def write_policy_from_config(config: Any) -> WritePolicy:
    """六键 → 唯一 ``WritePolicy`` 构造口。

    白名单未配置时回落 ``bot_download_dir/export``（＝今日导出腿唯一落点，现网
    零变更）；``external_verdict`` 显式接 ``sendable_verdict``——前席建议口径：
    正向注册名册那一半也在装配层接上，运行器缺省只兜禁触那一族，两层互不遮蔽。
    """
    configured = [
        str(item) for item in config.bot_files_write_allowed_dirs if str(item or "").strip()
    ]
    roots = configured or [str(Path(str(config.bot_download_dir)) / "export")]
    return policy_for_roots(
        roots,
        limits=WriteLimits(
            max_file_bytes=int(config.bot_files_write_max_bytes),
            daily_create_limit=int(config.bot_files_write_daily_create),
            daily_replace_limit=int(config.bot_files_write_daily_replace),
        ),
        external_verdict=sendable_verdict,
    )


def _read_policy_from_config(config: Any) -> WritePolicy:
    """回读侧策略：与写侧同根，只把单文件限额换成回读那一枚键。"""
    base = write_policy_from_config(config)
    return dataclass_replace(
        base,
        limits=dataclass_replace(
            base.limits,
            max_file_bytes=int(config.bot_files_read_confined_max_bytes),
        ),
    )


def adjudicate_file_write(
    actor_level: Any = None,
    actor_roles: Sequence[str] | None = None,
    *,
    internal_origin: bool = False,
) -> Any:
    """文件域落盘/外发动作的**唯一裁决出口**：向 ``safety_exec.policy.decide`` 问一次。

    这里不合成第二条判定——三态（``Permit`` / ``Deny`` / ``ConsentRequired``）原样
    交回，可信级、角色下限、档位分岔全在 ``safety_exec`` 一侧。缺省参数**不给
    内部发起**（``internal_origin=False``）：入站消息派生的动作必须由装配层申报
    可信级与角色，拿不到就是 ``Deny(untrusted_source)``，不许蒙。
    """
    return safety_policy.decide(
        action=ActionId.FILE_WRITE,
        actor_level=actor_level,
        actor_roles=actor_roles,
        internal_origin=internal_origin,
    )


def _adjudication_block(decision: Any) -> DocumentOpResult | None:
    """``Permit``→None（继续）；其余一律出拒绝句子，**绝不自批**。"""
    if isinstance(decision, safety_policy.Permit):
        return None
    if isinstance(decision, safety_policy.Deny):
        return DocumentOpResult(
            False,
            f"这一步没被放行：{decision.plain_text}",
            reason_code=f"adjudicate:{decision.kind.value}",
        )
    # ConsentRequired：``file.write`` 在册 R1（管理员本会话一次性确认号）。确认回路
    # 对非配置面动作今天还没接——那是第 18 项同意面的活（施工单二交主代理）。
    # 本口的处置＝如实说「要先过一次确认」并**一个字节都不写**，不自批、不放宽。
    tier = str(getattr(decision, "tier", "") or "R1")
    return DocumentOpResult(
        False,
        f"写文件要先过一次确认（{tier} 档：管理员在本会话给一次性确认号）。"
        "确认回路今天还没接到文件动作上——不自批、不写盘。",
        reason_code="consent_required",
    )


def _confirm_and_package(name: str, policy: WritePolicy, outcome: WriteOutcome) -> DocumentOpResult:
    """写结论 → 受限回读确认 → 出站件。回读对不上账就不认成功（防「谎报落盘」）。"""
    if outcome.denied:
        return DocumentOpResult(
            False,
            f"这次没写成：{outcome.error_message() or plain_reason(outcome.reason_code)}",
            reason_code=outcome.reason_code,
        )
    path = outcome.path
    if path is None:
        return DocumentOpResult(
            False, "写盘口没交回落点，这次不认成功。", reason_code="io_error"
        )
    read_back = read_confined_bytes(path.name, policy=policy)
    if isinstance(read_back, WriteOutcome):
        return DocumentOpResult(
            False,
            f"写完回读失败，不谎报送达：{read_back.error_message() or plain_reason(read_back.reason_code)}",
            reason_code="readback_failed",
        )
    data, digest = read_back
    if outcome.written_bytes and digest != outcome.sha256:
        return DocumentOpResult(
            False,
            "写完回读的指纹与落盘结论对不上，这次不认成功。",
            reason_code="readback_mismatch",
        )
    if outcome.written_bytes == 0:
        verb_text = "看过一遍：内容没变化，没动笔"
    elif outcome.verb == "replace":
        verb_text = "已改好并回读确认"
    else:
        verb_text = "已建好并回读确认"
    return DocumentOpResult(
        True,
        f"{verb_text}：{path.name}（{len(data)} 字节，sha {digest[:12]}）",
        file_parts=[{"file": str(path), "name": path.name}],
    )


def run_document_create(
    name: str,
    content: str | bytes,
    *,
    config: Any,
    decision: Any,
    ledger: DailyQuotaLedger | None = None,
    date_key: Callable[[], str] | None = None,
    staging_dir: Path | str | None = None,
) -> DocumentOpResult:
    """「创建文件」的装配入口：先过裁决，再落盘（fail-closed）。

    ``decision`` **必填**且应是 ``adjudicate_file_write`` 的产物——签名上不留
    「忘了问裁决口也能调」的形状；缺了它连调用都完不成，比注释约定硬。
    """
    if not bool(config.bot_files_write_enabled):
        return DocumentOpResult(False, _WRITE_DISABLED_REPLY, reason_code="write_disabled")
    blocked = _adjudication_block(decision)
    if blocked is not None:
        return blocked
    policy = write_policy_from_config(config)
    outcome = create_document_file(
        name, content, policy=policy, ledger=ledger, date_key=date_key,
        staging_dir=staging_dir,
    )
    return _confirm_and_package(name, policy, outcome)


def run_document_revise(
    name: str,
    transform: Callable[[bytes], object],
    *,
    config: Any,
    decision: Any,
    ledger: DailyQuotaLedger | None = None,
    date_key: Callable[[], str] | None = None,
    staging_dir: Path | str | None = None,
) -> DocumentOpResult:
    """「改这份文档」的装配入口：裁决 → 读-改-写单口 → 回读确认 → 出站件。

    transform 只收字节、拿不到路径（运行器教义）；目标必须在（``target_missing``
    绝不借改之名新建）；同字节不写、不烧配额。其余与 ``run_document_create`` 同形。
    """
    if not bool(config.bot_files_write_enabled):
        return DocumentOpResult(False, _WRITE_DISABLED_REPLY, reason_code="write_disabled")
    blocked = _adjudication_block(decision)
    if blocked is not None:
        return blocked
    policy = write_policy_from_config(config)
    outcome = revise_in_place(
        name, transform, policy=policy, ledger=ledger, date_key=date_key,
        staging_dir=staging_dir,
    )
    return _confirm_and_package(name, policy, outcome)


def read_document_file(name: str, *, config: Any) -> tuple[bytes, str] | DocumentOpResult:
    """受限回读的装配入口（「改这份」前半程：先看现状）。吃回读那一枚限额键。"""
    policy = _read_policy_from_config(config)
    found = read_confined_bytes(name, policy=policy)
    if isinstance(found, WriteOutcome):
        return DocumentOpResult(
            False,
            f"这份文件我按受限回读口径读不了：{found.error_message() or plain_reason(found.reason_code)}",
            reason_code=found.reason_code,
        )
    return found


def resolve_document_ref(name: str, *, config: Any) -> Path | DocumentOpResult:
    """写侧白名单内的寻址口（装配层给「改这份」定把手用；判据全在咽喉）。"""
    located = resolve_existing(name, policy=write_policy_from_config(config))
    if isinstance(located, WriteOutcome):
        return DocumentOpResult(
            False,
            f"这份文件我够不着：{located.error_message() or plain_reason(located.reason_code)}",
            reason_code=located.reason_code,
        )
    return located


def _export_docx(blocks: list[DocumentBlock], target: Path, title: str) -> tuple[Path, str]:
    import docx  # python-docx

    document = docx.Document()
    if title:
        document.add_heading(title, level=0)
    for block in blocks:
        if block.kind == "heading":
            document.add_heading(block.text, level=max(1, min(4, block.level)))
        elif block.kind == "bullet":
            document.add_paragraph(block.text, style="List Bullet")
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            run.font.name = "Courier New"
        elif block.kind == "table":
            table = document.add_table(rows=len(block.rows), cols=max(len(r) for r in block.rows))
            for row_index, row in enumerate(block.rows):
                for col_index, cell in enumerate(row):
                    table.rows[row_index].cells[col_index].text = cell
        else:
            document.add_paragraph(block.text)
    document.save(str(target))
    return target, ""


def _export_xlsx(blocks: list[DocumentBlock], target: Path) -> tuple[Path, str]:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    row_index = 1
    wrote_table = False
    for block in blocks:
        if block.kind == "table":
            for row in block.rows:
                for col_index, cell in enumerate(row, start=1):
                    sheet.cell(row=row_index, column=col_index, value=cell)
                row_index += 1
            wrote_table = True
            row_index += 1
        else:
            sheet.cell(row=row_index, column=1, value=block.text)
            row_index += 1
    if not wrote_table:
        sheet.title = "内容"
    workbook.save(str(target))
    return target, ""


def _export_pptx(blocks: list[DocumentBlock], target: Path, title: str) -> tuple[Path, str]:
    from pptx import Presentation
    from pptx.util import Pt

    presentation = Presentation()
    if title:
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = title
    current_bullets: list[str] = []
    current_heading = ""

    def _flush() -> None:
        nonlocal current_bullets, current_heading
        if not current_bullets:
            return
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = current_heading or title or "内容"
        body = slide.placeholders[1].text_frame
        for index, item in enumerate(current_bullets[:8]):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = item
            paragraph.font.size = Pt(18)
        current_bullets = []

    for block in blocks:
        if block.kind == "heading":
            _flush()
            current_heading = block.text
        elif block.kind == "bullet" or (
            block.kind == "para" and len(current_bullets) < 8
        ):
            current_bullets.append(block.text)
    _flush()
    if not presentation.slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = title or "空文档"
    presentation.save(str(target))
    return target, ""


def _export_pdf(blocks: list[DocumentBlock], target: Path, title: str) -> tuple[Path, str]:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    font_path = _find_cjk_font()
    if font_path:
        pdf.add_font("cjk", "", font_path)
        pdf.set_font("cjk", size=12)
    else:
        # 无中文字体时降级：仅保证不崩溃，报告由调用方提示。
        pdf.set_font("helvetica", size=12)

    def _line(height: float, text: str) -> None:
        # fpdf2 的 multi_cell 默认 new_x=RIGHT 不回左边距，必须手动复位，
        # 否则连续两行会因剩余宽度为 0 抛“无横向空间”。
        pdf.multi_cell(0, height, text)
        pdf.set_x(pdf.l_margin)

    if title:
        pdf.set_font_size(18)
        _line(10, title)
        pdf.set_font_size(12)
        pdf.ln(2)
    for block in blocks:
        if block.kind == "heading":
            pdf.set_font_size(12 + max(0, 5 - block.level) * 2)
            _line(8, block.text)
            pdf.set_font_size(12)
        elif block.kind == "bullet":
            # 用 ASCII 项目符号：中文字体缺 U+2022 字形时 fpdf2 会按零宽度抛错。
            _line(7, f"- {block.text}")
        elif block.kind == "code":
            for line in block.text.splitlines():
                _line(6, f"    {line}")
        elif block.kind == "table":
            for row in block.rows:
                _line(7, "  ".join(row))
        else:
            _line(7, block.text)
        pdf.ln(1)
    pdf.output(str(target))
    return target, ""
